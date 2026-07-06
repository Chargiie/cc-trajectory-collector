#!/usr/bin/env python3
"""从 xlsx 批量生产轨迹(2 路工作池,可断点续跑)。

- 从 xlsx 读「生产方 == OWNER 且未标记」的行作为工作队列;
- 一个共享代理(LOG_DIR 按 session 分流)+ 全局 baseurl 只设一次;
- N 个 worker 从队列取 query 跑(各自独立 workspace/TMPDIR),按 query 认领 session 文件;
- 每条产完:重建轨迹 -> 在 xlsx 写标记(已生产列=run目录、状态列)并存盘 -> 取下一条;
- 进度同时写 json sidecar(robust 续跑源)。中断后重跑自动跳过已标记行。

用法:
  TC_MODEL='claude-opus-4-8[1m]' python3 batch_xlsx.py --xlsx <path> --owner '@方飞腾' --workers 2
  加 --limit N 只做前 N 条(试跑)。
"""
import argparse, os, sys, json, time, socket, subprocess, shutil, signal, datetime, glob, threading, queue, re

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "lib"))
import config
from lib import build_trajectory, to_sft, render_html, pty_driver
from lib.validate import validate_run
from collect import slugify, free_port, sc_baseurl_get, sc_baseurl_set
import openpyxl

# 这些是本工具的跑批簿记列,写到空列 N/O/P,绝不碰人工维护的 K生产状态/L数据状态/M备注。
STATUS_COL = 14    # N「跑批状态」= 成功(...) / 失败:reason  ← 用户要的成功/失败
MARK_COL = 15      # O「run目录」= run 目录名(有值=本工具已跑过 → 续跑跳过)
TIME_COL = 16      # P「完成时间」

_xlsx_lock = threading.Lock()
_prog_lock = threading.Lock()


def ensure_headers(ws):
    if not ws.cell(1, STATUS_COL).value:
        ws.cell(1, STATUS_COL).value = "跑批状态"
    if not ws.cell(1, MARK_COL).value:
        ws.cell(1, MARK_COL).value = "run目录"
    if not ws.cell(1, TIME_COL).value:
        ws.cell(1, TIME_COL).value = "完成时间"


def load_jobs(ws, owner, done_ids):
    jobs = []
    for r in range(2, ws.max_row + 1):
        qid = ws.cell(r, 1).value
        q = ws.cell(r, 2).value
        if not qid or not q:
            continue
        if str(ws.cell(r, 10).value).strip() != owner:
            continue
        if ws.cell(r, MARK_COL).value or str(qid) in done_ids:
            continue   # 已标记 -> 跳过(续跑)
        jobs.append((r, str(qid), str(q)))
    return jobs


def mark_xlsx(xlsx, wb, ws, row, run_dir, status):
    with _xlsx_lock:
        ws.cell(row, MARK_COL).value = os.path.basename(run_dir) if run_dir else ""
        ws.cell(row, STATUS_COL).value = status
        ws.cell(row, TIME_COL).value = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        try:
            wb.save(xlsx)
        except Exception as e:
            sys.stderr.write(f"[warn] xlsx 写入失败(可能正被 Excel 打开): {e}\n")


def save_progress(prog_path, qid, rec):
    with _prog_lock:
        prog = {}
        if os.path.exists(prog_path):
            try: prog = json.load(open(prog_path))
            except Exception: prog = {}
        prog[qid] = rec
        json.dump(prog, open(prog_path, "w"), ensure_ascii=False, indent=2)


def worker(jobs_q, port, log_dir, args, xlsx, wb, ws, prog_path, stats):
    while True:
        try:
            row, qid, q = jobs_q.get_nowait()
        except queue.Empty:
            return
        run_tag = f"{datetime.datetime.now().strftime('%Y%m%d-%H%M%S')}-{qid.replace('_','-')}"
        run_dir = os.path.join(config.RUNS_DIR, f"{run_tag}_{slugify(q)}")
        workspace = os.path.join(run_dir, "workspace")
        tmp_dir = os.path.join(run_dir, "tmp")
        os.makedirs(workspace, exist_ok=True)
        os.makedirs(tmp_dir, exist_ok=True)
        print(f"[{qid}] 开跑: {q[:36]}")
        try:
            res = pty_driver.run_session(
                q, workspace, None, port=port, log_dir=log_dir,
                model=args.model, max_seconds=args.max_seconds, tmpdir=tmp_dir, log_cb=None,
            )
            if res.get("reason") in ("blocked_on_ask", "sse_truncated", "tool_hang"):
                raise RuntimeError(f"跳过({res['reason']}):卡在反问/截断/工具hang,已放行 worker")
            sf = res.get("session_file")
            if not (sf and os.path.exists(sf)):
                raise RuntimeError(f"未认领到 session 文件 ({res.get('reason')})")
            calls_path = os.path.join(run_dir, "calls.jsonl")
            shutil.copy(sf, calls_path)
            traj = build_trajectory.build(calls_path)
            json.dump(traj, open(os.path.join(run_dir, "trajectory.json"), "w"), ensure_ascii=False, indent=2)
            sft = to_sft.to_sft(traj, session_id=qid)
            json.dump([sft], open(os.path.join(run_dir, "trajectory_sft.json"), "w"), ensure_ascii=False, indent=2)
            if not args.no_html:
                open(os.path.join(run_dir, "trajectory.html"), "w").write(render_html.render(traj))
            # 即时校验:检测产物质量问题并 tag
            warnings = validate_run(run_dir, sft_data=sft)
            warn_tag = f" ⚠{'|'.join(warnings)}" if warnings else ""
            status = f"成功({res['reason']},{traj['meta']['num_main_calls']}步,{traj['meta'].get('tokens',{}).get('total_tokens',0)}tok){warn_tag}"
            mark_xlsx(xlsx, wb, ws, row, run_dir, status)
            save_progress(prog_path, qid, {"status": "done", "run_dir": run_dir, "reason": res["reason"],
                                           "main": traj["meta"]["num_main_calls"], "model": traj["meta"]["model"],
                                           "tokens": traj["meta"].get("tokens", {}),
                                           "warnings": warnings})
            stats["done"] += 1
            print(f"[{qid}] ✓ {status} -> {os.path.basename(run_dir)} | 累计完成 {stats['done']}")
        except Exception as e:
            mark_xlsx(xlsx, wb, ws, row, run_dir, f"失败:{e}")
            save_progress(prog_path, qid, {"status": "fail", "run_dir": run_dir, "error": str(e)})
            stats["fail"] += 1
            print(f"[{qid}] ✗ 失败: {e}")
        finally:
            for md in glob.glob(os.path.join(config.HOME, ".claude", "projects", f"*{run_tag}*")):
                shutil.rmtree(md, ignore_errors=True)
            jobs_q.task_done()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--owner", default="@方飞腾")
    ap.add_argument("--workers", type=int, default=2)
    ap.add_argument("--model", default=config.MODEL)
    ap.add_argument("--max-seconds", type=int, default=config.MAX_SECONDS)
    ap.add_argument("--no-html", action="store_true")
    ap.add_argument("--limit", type=int, default=0, help=">0 时只做前 N 条(试跑)")
    args = ap.parse_args()

    xlsx = args.xlsx
    prog_path = os.path.join(config.RUNS_DIR, "_batch_progress.json")
    os.makedirs(config.RUNS_DIR, exist_ok=True)
    done_ids = set()
    if os.path.exists(prog_path):
        try:
            done_ids = {k for k, v in json.load(open(prog_path)).items() if v.get("status") == "done"}
        except Exception:
            pass

    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    ensure_headers(ws)
    jobs = load_jobs(ws, args.owner, done_ids)
    if args.limit > 0:
        jobs = jobs[:args.limit]
    print(f"=== 批量生产 ===")
    print(f"负责人={args.owner} | 模型={args.model} | 并发={args.workers} | 待做 {len(jobs)} 条(已跳过已完成)")
    if not jobs:
        print("没有待做的 query(都标记完成了?)"); return

    ts = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    for pat in ("/tmp/pg-*.png", "/tmp/checkedges_*", "/tmp/checklayout_*"):
        for p in glob.glob(pat):
            try: shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)
            except OSError: pass

    port = free_port(config.PROXY_PORT)
    log_dir = os.path.join(config.RUNS_DIR, f"_sharedlog_{ts}")
    os.makedirs(log_dir, exist_ok=True)
    proxy_env = dict(os.environ)
    proxy_env.update({"PROXY_PORT": str(port), "UPSTREAM_URL": config.UPSTREAM_URL, "LOG_DIR": log_dir})
    if config.THINKING_HACK:
        proxy_env["THINKING_HACK"] = "1"
    proxy_out = open(os.path.join(log_dir, "proxy.out"), "w")
    proxy = subprocess.Popen(["node", config.PROXY_JS], env=proxy_env, stdout=proxy_out, stderr=subprocess.STDOUT)
    for _ in range(40):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(("127.0.0.1", port)) == 0: break
        time.sleep(0.1)
    print(f"[proxy] 共享端口 {port} | 日志 {log_dir}")

    jobs_q = queue.Queue()
    for j in jobs:
        jobs_q.put(j)
    stats = {"done": 0, "fail": 0}
    orig = sc_baseurl_get()
    print(f"[baseurl] {orig} -> http://127.0.0.1:{port}(只设一次)")
    try:
        sc_baseurl_set(f"http://127.0.0.1:{port}")
        threads = []
        for i in range(args.workers):
            t = threading.Thread(target=worker, args=(jobs_q, port, log_dir, args, xlsx, wb, ws, prog_path, stats))
            t.start(); threads.append(t); time.sleep(3)
        for t in threads:
            t.join()
    finally:
        if orig:
            sc_baseurl_set(orig); print(f"[baseurl] 已还原 -> {orig}")
        try: proxy.send_signal(signal.SIGTERM); proxy.wait(timeout=5)
        except Exception: proxy.kill()
        proxy_out.close()

    print(f"\n=== 本轮结束 === 完成 {stats['done']} | 失败 {stats['fail']} | xlsx 标记已写入 {xlsx}")


if __name__ == "__main__":
    main()
